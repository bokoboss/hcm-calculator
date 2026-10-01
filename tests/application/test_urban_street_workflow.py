import base64
import csv
from copy import deepcopy
import json
from io import BytesIO, StringIO
from pathlib import Path
import re

import pytest
from openpyxl import load_workbook

from hcmcalc.application.project import (
    ProjectFileError,
    compare_scenarios,
    duplicate_scenario,
    load_project,
    project_to_json,
    record_result,
    save_analysis_to_project,
    update_scenario_inputs,
)
from hcmcalc.application.registry import get_analysis_definition, list_analysis_definitions
from hcmcalc.application.workflows import (
    ApplicationWorkflowError,
    StaleResultError,
    export_current_workflow,
    normalized_workflow_inputs,
    workflow_for_method,
)
from hcmcalc.ui.units import FEET_TO_METERS, MILES_TO_KILOMETERS
import hcmcalc.urban_street_ch18 as ch18


METHOD_ID = "urban_street_segment"
TEMPLATE_ID = "USS-CH30-EP1"


def _chapter18_snapshot(unit_system="imperial", displayed=None):
    workflow = workflow_for_method(METHOD_ID)
    values = displayed or workflow.starting_values(TEMPLATE_ID, unit_system)["displayed_inputs"]
    return workflow.calculate(template_id=TEMPLATE_ID, unit_system=unit_system, displayed_inputs=values)


def _assert_equivalent(left, right, path="value"):
    if isinstance(left, bool) or isinstance(right, bool):
        assert left is right, path
    elif isinstance(left, (int, float)) and isinstance(right, (int, float)):
        assert left == pytest.approx(right, rel=0, abs=1e-12), path
    elif isinstance(left, dict) and isinstance(right, dict):
        assert set(left) == set(right), path
        for key in left:
            _assert_equivalent(left[key], right[key], f"{path}.{key}")
    elif isinstance(left, list) and isinstance(right, list):
        assert len(left) == len(right), path
        for index, (left_item, right_item) in enumerate(zip(left, right, strict=True)):
            _assert_equivalent(left_item, right_item, f"{path}[{index}]")
    else:
        assert left == right, path


def test_registry_adds_frozen_chapter_18_identity_without_changing_existing_seven():
    definition = get_analysis_definition(METHOD_ID)
    assert definition is not None
    assert {
        "method_id": definition.method_id,
        "family": definition.family,
        "method_identifier": definition.method_identifier,
        "engine_method_identifier": definition.engine_method_identifier,
        "method_version": definition.method_version,
        "input_contract": definition.input_contract,
        "project_type": definition.project_type,
        "hcm_edition": definition.hcm_edition,
        "hcm_chapter": definition.hcm_chapter,
        "chapter_reference": definition.chapter_reference,
        "supported_unit_systems": definition.supported_unit_systems,
        "availability": definition.availability,
    } == {
        "method_id": "urban_street_segment",
        "family": "urban_streets",
        "method_identifier": "hcm7_urban_street_segment",
        "engine_method_identifier": "urban_street_segment_ch18_v0_1",
        "method_version": "hcm_7_0_bounded_v1",
        "input_contract": "hcm7_ch18_bounded_signalized_15min_rht_reference_v1",
        "project_type": "manual_urban_street_segment_v1",
        "hcm_edition": "HCM 7.0",
        "hcm_chapter": "18",
        "chapter_reference": "HCM 7.0 Chapter 18; Chapter 30 Example Problem 1",
        "supported_unit_systems": ("metric", "imperial"),
        "availability": "qualified_bounded",
    }
    assert {item.method_id for item in list_analysis_definitions()} == {
        "two_lane_segment",
        "two_lane_facility",
        "multilane_segment",
        "basic_freeway_segment",
        "weaving_segment",
        "merge_segment",
        "diverge_segment",
        METHOD_ID,
    }
    assert definition.capabilities == ("segment", "audit", "external_through_performance")
    assert definition.scope_summary_keys == (
        "method.urban_street_segment.scope.bounded_signalized_15min",
        "method.urban_street_segment.scope.rht_reference",
    )


def test_existing_seven_application_identities_are_frozen():
    expected = {
        "two_lane_segment": ("hcm7_two_lane_highway_segment", "hcm7_ch15_two_lane_motorized", "hcm7.0", "phase_5_product_integration", "manual_single_segment"),
        "two_lane_facility": ("hcm7_two_lane_highway_facility", "hcm7_ch15_two_lane_motorized", "hcm7.0", "phase_5_product_integration", "manual_two_lane_facility_v1"),
        "multilane_segment": ("hcm7_multilane_los", "hcm7_multilane_los", "v0.1", "phase_8", "manual_multilane_v0"),
        "basic_freeway_segment": ("hcm7_basic_freeway_segment", "hcm7_basic_freeway_segment", "phase_9_engine", "phase_10_product_integration", "manual_basic_freeway_v0"),
        "weaving_segment": ("weaving_segment", "hcm7_v70_freeway_weaving_segment", "hcm_7_0", "hcm_7_0_weaving_segment_operational_v1", "manual_freeway_weaving_segment_v1"),
        "merge_segment": ("merge_segment", "hcm7_v70_freeway_merge_segment", "hcm_7_0", "hcm7_v70_chapter_14_isolated_right_side_one_lane_merge_operational", "manual_freeway_merge_segment_v1"),
        "diverge_segment": ("diverge_segment", "hcm7_v70_freeway_diverge_segment", "hcm_7_0", "hcm7_v70_chapter_14_isolated_right_side_one_lane_diverge_operational", "manual_freeway_diverge_segment_v1"),
    }
    for method_id, identity in expected.items():
        definition = get_analysis_definition(method_id)
        assert (definition.method_identifier, definition.engine_method_identifier, definition.method_version,
                definition.input_contract, definition.project_type) == identity


def test_workflow_dispatch_and_starters_are_available():
    workflow = workflow_for_method(METHOD_ID)
    assert {item["template_id"] for item in workflow.templates()["templates"]} == {
        TEMPLATE_ID,
        "blank_custom",
    }
    assert workflow.starting_values(TEMPLATE_ID, "imperial")["validation_status"] == "reference_fixture"


def test_public_field_metadata_uses_text_and_bounded_choice_types():
    fields = {field["key"]: field for field in workflow_for_method(METHOD_ID).templates()["fields"]}
    for key in (
        "subject_direction", "through_movement_id", "external_source_class",
        "external_source_tool", "external_source_method_note",
        "external_hcm_edition_note", "external_direction",
        "external_through_movement_id", "external_scenario_note",
    ):
        assert fields[key]["kind"] == "text"
    for key in ("control_type", "external_control_type"):
        assert fields[key]["kind"] == "choice"
        assert fields[key]["options"] == ["signalized"]
    assert fields["segment_length"]["kind"] == "number"
    assert fields["v_th_veh_h"]["kind"] == "number"
    assert fields["through_lane_count"]["kind"] == "integer"
    assert fields["external_analysis_period_min"]["kind"] == "integer"
    assert fields["demand_balanced"]["kind"] == "boolean"
    assert fields["spillback_present"]["kind"] == "boolean"
    assert fields["access_point_delays_s_veh"]["kind"] == "number_list"


def test_advertised_chapter18_scope_keys_are_translated_in_both_catalogs():
    definition = get_analysis_definition(METHOD_ID)
    catalog_path = Path(__file__).resolve().parents[2] / "frontend" / "src" / "i18n" / "catalog.ts"
    source = catalog_path.read_text(encoding="utf-8")
    english = source.split("  en: {", 1)[1].split("  th: {", 1)[0]
    thai = source.split("  th: {", 1)[1].rsplit("\n};", 1)[0]

    for key in definition.scope_summary_keys:
        for section in (english, thai):
            match = re.search(rf"^\s*'{re.escape(key)}':\s*'([^']*)',?$", section, re.MULTILINE)
            assert match is not None, key
            assert match.group(1) and match.group(1) != key


@pytest.mark.parametrize("field", ["segment_length", "posted_speed_limit"])
def test_invalid_scaled_displayed_values_return_structured_application_errors(field):
    workflow = workflow_for_method(METHOD_ID)
    displayed = workflow.starting_values(TEMPLATE_ID, "imperial")["displayed_inputs"]
    displayed[field] = "abc"

    validation = workflow.validate(template_id=TEMPLATE_ID, unit_system="imperial", displayed_inputs=displayed)
    assert validation["valid"] is False
    assert validation["ready"] is False
    assert validation["errors"][0]["code"] == "invalid_input"
    assert validation["errors"][0]["field"] == field
    with pytest.raises(ApplicationWorkflowError) as error:
        workflow.calculate(template_id=TEMPLATE_ID, unit_system="imperial", displayed_inputs=displayed)
    assert error.value.code == "invalid_input"
    assert error.value.details == {"field": field}


def test_validation_rejects_eq18_6_domain_before_calculation(monkeypatch):
    workflow = workflow_for_method(METHOD_ID)
    displayed = workflow.starting_values(TEMPLATE_ID, "imperial")["displayed_inputs"]
    displayed["v_m_veh_h"] = 5000
    monkeypatch.setattr(ch18.UrbanStreetSegmentMethod, "calculate", lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("validation calculated")))

    validation = workflow.validate(template_id=TEMPLATE_ID, unit_system="imperial", displayed_inputs=displayed)
    assert validation["valid"] is False
    assert validation["ready"] is False
    assert validation["errors"][0]["code"] == "invalid_input"
    assert validation["errors"][0]["field"] is None


def test_validation_and_calculation_reject_same_below_25_bffs_starter_case():
    workflow = workflow_for_method(METHOD_ID)
    displayed = workflow.starting_values(TEMPLATE_ID, "imperial")["displayed_inputs"]
    displayed["s_calib"] = -15.77965142857143

    validation = workflow.validate(
        template_id=TEMPLATE_ID,
        unit_system="imperial",
        displayed_inputs=displayed,
    )
    assert validation["valid"] is False
    assert validation["ready"] is False
    assert validation["errors"][0]["code"] == "invalid_input"

    with pytest.raises(ApplicationWorkflowError) as error:
        workflow.calculate(
            template_id=TEMPLATE_ID,
            unit_system="imperial",
            displayed_inputs=displayed,
        )
    assert error.value.code == "invalid_input"


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("control_type", "two_way_stop"),
        ("analysis_period_min", 60),
        ("demand_balanced", False),
        ("demand_adjustments_resolved", False),
        ("capacity_effects_resolved", False),
        ("spillback_present", True),
        ("access_point_delays_s_veh", [-0.1, 0.2]),
        ("upstream_intersection_width", 1800),
    ],
)
def test_validation_errors_identify_directly_offending_chapter18_display_field(field, value):
    workflow = workflow_for_method(METHOD_ID)
    displayed = workflow.starting_values(TEMPLATE_ID, "imperial")["displayed_inputs"]
    displayed[field] = value

    validation = workflow.validate(template_id=TEMPLATE_ID, unit_system="imperial", displayed_inputs=displayed)

    assert validation["valid"] is False
    assert validation["ready"] is False
    assert validation["errors"][0]["code"] == "invalid_input"
    assert validation["errors"][0]["field"] == field


def test_chapter18_project_comparison_reports_canonical_current_output_deltas_without_rerun(monkeypatch):
    base = _chapter18_snapshot()
    project = save_analysis_to_project(base, project_name="Chapter 18 comparison")
    analysis = project["analyses"][0]
    left = analysis["scenarios"][0]
    duplicated = duplicate_scenario(
        project,
        analysis_id=analysis["analysis_id"],
        scenario_id=left["scenario_id"],
        scenario_name="Curb alternative",
    )
    right = duplicated["analyses"][0]["scenarios"][1]
    changed_inputs = deepcopy(base["displayed_inputs"])
    changed_inputs["curb_proportion"] = 0.71
    changed = _chapter18_snapshot(displayed=changed_inputs)
    current = record_result(
        duplicated,
        analysis_id=analysis["analysis_id"],
        scenario_id=right["scenario_id"],
        snapshot=changed,
    )

    assert left["result_status"] == "current"
    assert current["analyses"][0]["scenarios"][1]["result_status"] == "current"
    assert left["result"]["engine_result"]["outputs"]["level_of_service"] == current["analyses"][0]["scenarios"][1]["result"]["engine_result"]["outputs"]["level_of_service"]

    monkeypatch.setattr(ch18.UrbanStreetSegmentMethod, "calculate", lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("comparison reran engine")))
    comparison = compare_scenarios(
        current,
        analysis_id=analysis["analysis_id"],
        left_scenario_id=left["scenario_id"],
        right_scenario_id=right["scenario_id"],
    )
    deltas = {delta["key"]: delta for delta in comparison["numeric_deltas"]}
    assert comparison["los_grade_transition"]["changed"] is False
    assert comparison["recalculated"] is False
    assert "travel_speed_mph" in deltas
    assert deltas["travel_speed_mph"]["left"] != deltas["travel_speed_mph"]["right"]


def test_project_canonicalization_recovers_chapter_18_inputs_without_calculation():
    workflow = workflow_for_method(METHOD_ID)
    displayed = workflow.starting_values(TEMPLATE_ID, "imperial")["displayed_inputs"]
    normalized = normalized_workflow_inputs(
        METHOD_ID,
        template_id=TEMPLATE_ID,
        unit_system="imperial",
        displayed_inputs=displayed,
    )
    assert normalized["v_m_veh_h"] == 1150
    assert normalized["external_through"]["v_th_veh_h"] == 968


def test_metric_and_imperial_chapter_30_starters_normalize_to_same_engine_inputs():
    workflow = workflow_for_method(METHOD_ID)
    metric = workflow.starting_values(TEMPLATE_ID, "metric")["displayed_inputs"]
    imperial = workflow.starting_values(TEMPLATE_ID, "imperial")["displayed_inputs"]
    assert metric["segment_length"] == 1800 * FEET_TO_METERS
    assert metric["posted_speed_limit"] == 35 * MILES_TO_KILOMETERS

    metric_native = normalized_workflow_inputs(
        METHOD_ID,
        template_id=TEMPLATE_ID,
        unit_system="metric",
        displayed_inputs=metric,
    )
    imperial_native = normalized_workflow_inputs(
        METHOD_ID,
        template_id=TEMPLATE_ID,
        unit_system="imperial",
        displayed_inputs=imperial,
    )
    _assert_equivalent(metric_native, imperial_native, "normalized")
    metric_result = workflow.calculate(template_id=TEMPLATE_ID, unit_system="metric", displayed_inputs=metric)
    imperial_result = workflow.calculate(template_id=TEMPLATE_ID, unit_system="imperial", displayed_inputs=imperial)
    _assert_equivalent(metric_result["result"], imperial_result["result"], "result")
    assert metric_result["calculation_fingerprint"] == workflow.calculate(template_id=TEMPLATE_ID, unit_system="metric", displayed_inputs=metric)["calculation_fingerprint"]
    assert metric_result["input_snapshot_fingerprint"] != imperial_result["input_snapshot_fingerprint"]


def test_blank_is_fail_closed_and_invalid_external_preconditions_are_structured():
    workflow = workflow_for_method(METHOD_ID)
    blank = workflow.starting_values("blank_custom", "imperial")["displayed_inputs"]
    assert all(blank[key] is None for key in (
        "demand_balanced", "demand_adjustments_resolved", "capacity_effects_resolved",
        "spillback_present", "external_source_class", "external_source_tool",
        "external_source_method_note", "external_hcm_edition_note", "external_direction",
        "external_control_type", "external_through_movement_id", "external_scenario_note",
    ))
    assert workflow.validate(template_id="blank_custom", unit_system="imperial", displayed_inputs=blank)["ready"] is False
    valid = workflow.starting_values(TEMPLATE_ID, "imperial")["displayed_inputs"]
    assert workflow.validate(template_id=TEMPLATE_ID, unit_system="imperial", displayed_inputs=valid)["ready"] is True
    for field, value in (("external_source_tool", None), ("control_type", "two_way_stop"),
                         ("analysis_period_min", 60), ("spillback_present", True),
                         ("demand_balanced", False)):
        changed = deepcopy(valid)
        changed[field] = value
        validation = workflow.validate(template_id=TEMPLATE_ID, unit_system="imperial", displayed_inputs=changed)
        assert validation["ready"] is False
        assert validation["errors"][0]["code"] == "invalid_input"
        assert "message" in validation["errors"][0]


def test_chapter_30_calculation_uses_engine_once_and_preserves_v_m_separately(monkeypatch):
    original = ch18.UrbanStreetSegmentMethod.calculate
    calls = []

    def counted(self, inputs):
        calls.append(inputs)
        return original(self, inputs)

    monkeypatch.setattr(ch18.UrbanStreetSegmentMethod, "calculate", counted)
    snapshot = _chapter18_snapshot()
    out = snapshot["result"]["outputs"]
    assert len(calls) == 1
    assert snapshot["result"]["method"] == "urban_street_segment_ch18_v0_1"
    assert snapshot["result"]["result_contract_version"] == "hcm7_ch18_bounded_signalized_15min_rht_reference_v1"
    assert snapshot["normalized_inputs"]["v_m_veh_h"] == 1150
    assert snapshot["normalized_inputs"]["external_through"]["v_th_veh_h"] == 968
    assert out["s0_mph"] == pytest.approx(42.05)
    assert out["f_cs_mph"] == pytest.approx(-0.329)
    assert out["f_a_mph"] == pytest.approx(-0.941349)
    assert out["f_pk_mph"] == 0
    assert out["base_free_flow_speed_mph"] == pytest.approx(40.779651)
    assert out["f_l"] == pytest.approx(0.964436)
    assert out["free_flow_speed_mph"] == pytest.approx(39.329383)
    assert out["f_v"] == pytest.approx(1.034028)
    assert out["total_access_delay_s_veh"] == pytest.approx(0.387)
    assert out["running_time_s"] == pytest.approx(33.542718)
    assert out["running_speed_mph"] == pytest.approx(36.588350)
    assert out["travel_speed_mph"] == pytest.approx(23.668436)
    assert out["through_v_c"] == pytest.approx(968 / 1848)
    assert out["level_of_service"] == "C"
    assert "right-hand-traffic reference" in " ".join(snapshot["audit"]["assumptions"])
    assert snapshot["audit"]["external_through_provenance"] == out["external_through_provenance"]
    assert snapshot["presentation"]["metrics"][0]["value"] == out["travel_speed_mph"]


def test_v_over_c_above_one_forces_los_f_without_discarding_chapter18_speed():
    workflow = workflow_for_method(METHOD_ID)
    inputs = workflow.starting_values(TEMPLATE_ID, "imperial")["displayed_inputs"]
    inputs["v_th_veh_h"] = 1900
    snapshot = workflow.calculate(template_id=TEMPLATE_ID, unit_system="imperial", displayed_inputs=inputs)
    assert snapshot["result"]["outputs"]["through_v_c"] > 1
    assert snapshot["result"]["outputs"]["level_of_service"] == "F"
    assert snapshot["result"]["outputs"]["travel_speed_mph"] is not None
    assert snapshot["calculation_state"]["presentation_state"] == "valid_current_result"


def test_project_v2_roundtrip_retains_result_without_engine_rerun_and_edit_stales(monkeypatch):
    snapshot = _chapter18_snapshot()
    project = save_analysis_to_project(snapshot, project_name="Chapter 18 study")
    analysis = project["analyses"][0]
    scenario = analysis["scenarios"][0]
    assert project["schema_version"] == "2.0"
    assert analysis["method_id"] == METHOD_ID
    assert scenario["result_status"] == "current"

    def fail(*args, **kwargs):
        raise AssertionError("Project load must not rerun the engine")

    with monkeypatch.context() as patcher:
        patcher.setattr(ch18.UrbanStreetSegmentMethod, "calculate", fail)
        loaded = load_project(project_to_json(project))
    reopened = loaded["analyses"][0]["scenarios"][0]
    assert reopened["result_status"] == "current"
    assert reopened["result"]["engine_result"]["method"] == "urban_street_segment_ch18_v0_1"
    assert reopened["result"]["audit"]["external_through_provenance"] == snapshot["audit"]["external_through_provenance"]

    changed_inputs = deepcopy(snapshot["displayed_inputs"])
    changed_inputs["curb_proportion"] = 0.71
    changed = workflow_for_method(METHOD_ID).calculate(template_id=TEMPLATE_ID, unit_system="imperial", displayed_inputs=changed_inputs)
    stale = update_scenario_inputs(loaded, analysis_id=analysis["analysis_id"], scenario_id=scenario["scenario_id"], snapshot=changed)
    stale_scenario = stale["analyses"][0]["scenarios"][0]
    assert stale_scenario["result_status"] == "stale"
    assert stale_scenario["result"] is None
    current = record_result(stale, analysis_id=analysis["analysis_id"], scenario_id=scenario["scenario_id"], snapshot=changed)
    assert current["analyses"][0]["scenarios"][0]["result_status"] == "current"


@pytest.mark.parametrize("tamper", ("method_identifier", "engine_method_identifier", "input_contract", "calculation_fingerprint", "engine_result", "engine_result_method", "result_contract"))
def test_project_v2_rejects_chapter18_identity_and_result_tampering(tamper):
    project = save_analysis_to_project(_chapter18_snapshot())
    broken = deepcopy(project)
    analysis = broken["analyses"][0]
    scenario = analysis["scenarios"][0]
    if tamper in {"method_identifier", "engine_method_identifier", "input_contract"}:
        analysis[tamper] = "tampered"
    elif tamper == "calculation_fingerprint":
        scenario["calculation_fingerprint"] = "tampered"
    elif tamper == "engine_result":
        scenario["result"]["engine_result"]["outputs"]["calculation_type"] = "wrong_engine"
    elif tamper == "engine_result_method":
        scenario["result"]["engine_result"]["method"] = "wrong_engine"
    else:
        scenario["result"]["engine_result"]["result_contract_version"] = "wrong_contract"
    with pytest.raises(ProjectFileError):
        load_project(broken)


@pytest.mark.parametrize("export_format", ("csv", "xlsx", "markdown", "json"))
def test_current_result_exports_without_rerunning_and_preserves_external_provenance(monkeypatch, export_format):
    snapshot = _chapter18_snapshot()
    monkeypatch.setattr(ch18.UrbanStreetSegmentMethod, "calculate", lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("export reran engine")))
    exported = export_current_workflow(
        METHOD_ID, template_id=TEMPLATE_ID, unit_system="imperial",
        displayed_inputs=snapshot["displayed_inputs"],
        calculation_fingerprint=snapshot["calculation_fingerprint"],
        input_snapshot_fingerprint=snapshot["input_snapshot_fingerprint"],
        result=snapshot["result"], export_format=export_format,
    )
    assert exported["recalculated"] is False
    assert exported["calculation_fingerprint"] == snapshot["calculation_fingerprint"]
    assert exported["content"] or exported["content_base64"]
    if export_format == "json":
        assert "external_through_provenance" in exported["content"]


@pytest.mark.parametrize("unit_system", ("metric", "imperial"))
@pytest.mark.parametrize("export_format", ("json", "markdown", "csv", "xlsx"))
def test_chapter18_reports_keep_displayed_and_normalized_sections_in_all_formats(
    monkeypatch, unit_system, export_format
):
    workflow = workflow_for_method(METHOD_ID)
    snapshot = workflow.calculate(
        template_id=TEMPLATE_ID,
        unit_system=unit_system,
        displayed_inputs=workflow.starting_values(TEMPLATE_ID, unit_system)["displayed_inputs"],
    )
    assert snapshot["result"]["outputs"]["input_summary"] == snapshot["normalized_inputs"]
    original_displayed = deepcopy(snapshot["displayed_inputs"])
    monkeypatch.setattr(
        ch18.UrbanStreetSegmentMethod,
        "calculate",
        lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("report reran engine")),
    )
    exported = export_current_workflow(
        METHOD_ID,
        template_id=TEMPLATE_ID,
        unit_system=unit_system,
        displayed_inputs=snapshot["displayed_inputs"],
        calculation_fingerprint=snapshot["calculation_fingerprint"],
        input_snapshot_fingerprint=snapshot["input_snapshot_fingerprint"],
        result=snapshot["result"],
        export_format=export_format,
    )

    length_unit, speed_unit = ("m", "km/h") if unit_system == "metric" else ("ft", "mi/h")
    displayed_length = 548.64 if unit_system == "metric" else 1800
    displayed_speed = 56.327040000000004 if unit_system == "metric" else 35
    displayed_length_text = str(snapshot["displayed_inputs"]["segment_length"])
    displayed_speed_text = str(snapshot["displayed_inputs"]["posted_speed_limit"])
    normalized_length = snapshot["result"]["outputs"]["input_summary"]["segment_length_ft"]
    if export_format == "json":
        report = json.loads(exported["content"])
        assert "inputs_summary" in report and "normalized_engine_inputs_summary" in report
        displayed = {item["label"]: item for item in report["inputs_summary"]}
        normalized = {item["label"]: item for item in report["normalized_engine_inputs_summary"]}
        assert displayed["segment_length"] == {"label": "segment_length", "value": displayed_length, "unit": length_unit}
        assert displayed["posted_speed_limit"]["value"] == pytest.approx(displayed_speed)
        assert displayed["posted_speed_limit"]["unit"] == speed_unit
        assert normalized["segment_length_ft"]["value"] == pytest.approx(1800)
        assert normalized["segment_length_ft"]["unit"] == "ft (HCM-native)"
        if unit_system == "metric":
            assert displayed["segment_length"]["value"] == pytest.approx(548.64)
            assert displayed["segment_length"]["unit"] == "m"
            assert displayed["upstream_intersection_width"]["value"] == pytest.approx(15.24)
            assert displayed["upstream_intersection_width"]["unit"] == "m"
            assert displayed["signal_control_spacing"]["value"] == pytest.approx(548.64)
            assert displayed["signal_control_spacing"]["unit"] == "m"
            assert normalized["upstream_intersection_width_ft"]["value"] == pytest.approx(50)
            assert normalized["signal_control_spacing_ft"]["value"] == pytest.approx(1800)
            assert normalized["posted_speed_limit_mph"]["value"] == pytest.approx(35)
            assert normalized["posted_speed_limit_mph"]["unit"] == "mi/h (HCM-native)"
            assert displayed["external_source_tool"]["value"] == "HCM 7 Chapter 30 Example Problem 1"
            assert normalized["external_through.source_tool"]["value"] == "HCM 7 Chapter 30 Example Problem 1"
        assert set(displayed) != set(normalized)
    elif export_format == "markdown":
        text = exported["content"]
        assert "## Key Inputs" in text and "## Normalized Engine Inputs" in text
        assert f"| segment_length | {displayed_length_text} | {length_unit} |" in text
        assert f"| posted_speed_limit | {displayed_speed_text} | {speed_unit} |" in text
        assert f"| segment_length_ft | {normalized_length} | ft (HCM-native) |" in text
    elif export_format == "csv":
        rows = list(csv.reader(StringIO(exported["content"])))
        assert ["Inputs"] in rows and ["Normalized Engine Inputs"] in rows
        assert ["segment_length", displayed_length_text, length_unit] in rows
        assert ["posted_speed_limit", displayed_speed_text, speed_unit] in rows
        assert ["segment_length_ft", str(normalized_length), "ft (HCM-native)"] in rows
    else:
        workbook = load_workbook(BytesIO(base64.b64decode(exported["content_base64"])), data_only=False)
        rows = list(workbook["Inputs"].values)
        displayed_rows = {row[0]: row for row in rows if row[0] in {"segment_length", "posted_speed_limit"}}
        assert displayed_rows["segment_length"][1] == pytest.approx(displayed_length)
        assert displayed_rows["segment_length"][2] == length_unit
        assert displayed_rows["posted_speed_limit"][1] == pytest.approx(displayed_speed)
        assert displayed_rows["posted_speed_limit"][2] == speed_unit
        normalized_heading = rows.index(("Normalized Engine Inputs", None, None))
        assert rows[normalized_heading + 1] == ("Label", "Value", "Unit")
        normalized_rows = rows[normalized_heading + 2:]
        engine_length = next(row for row in normalized_rows if row[0] == "segment_length_ft")
        assert engine_length[1] == pytest.approx(1800)
        assert engine_length[2] == "ft (HCM-native)"

    assert exported["recalculated"] is False
    assert exported["calculation_fingerprint"] == snapshot["calculation_fingerprint"]
    revalidated = workflow.validate(
        template_id=TEMPLATE_ID,
        unit_system=unit_system,
        displayed_inputs=snapshot["displayed_inputs"],
    )
    assert revalidated["calculation_fingerprint"] == snapshot["calculation_fingerprint"]
    assert revalidated["input_snapshot_fingerprint"] == snapshot["input_snapshot_fingerprint"]
    assert snapshot["displayed_inputs"] == original_displayed


def test_spreadsheet_exports_neutralize_chapter18_provenance_without_mutating_source_or_other_exports():
    workflow = workflow_for_method(METHOD_ID)
    displayed = workflow.starting_values(TEMPLATE_ID, "imperial")["displayed_inputs"]
    dangerous = {
        "external_source_tool": '=WEBSERVICE("https://example.invalid")',
        "external_source_method_note": "+SUM(1,1)",
        "external_hcm_edition_note": "-1+1",
        "external_scenario_note": "@SUM(A1:A2)",
        "external_source_class": " \t=1+1",
    }
    displayed.update(dangerous)
    original = deepcopy(displayed)
    snapshot = workflow.calculate(template_id=TEMPLATE_ID, unit_system="imperial", displayed_inputs=displayed)
    source_fingerprint = snapshot["calculation_fingerprint"]

    def export(format_name):
        return export_current_workflow(
            METHOD_ID,
            template_id=TEMPLATE_ID,
            unit_system="imperial",
            displayed_inputs=displayed,
            calculation_fingerprint=snapshot["calculation_fingerprint"],
            input_snapshot_fingerprint=snapshot["input_snapshot_fingerprint"],
            result=snapshot["result"],
            export_format=format_name,
        )

    csv_output = export("csv")
    csv_cells = [cell for row in csv.reader(StringIO(csv_output["content"])) for cell in row]
    for value in dangerous.values():
        assert "'" + value in csv_cells

    xlsx_output = export("xlsx")
    workbook = load_workbook(BytesIO(base64.b64decode(xlsx_output["content_base64"])), data_only=False)
    cells = [cell for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row]
    for value in dangerous.values():
        literal = next(cell for cell in cells if cell.value == "'" + value)
        assert literal.data_type == "s"
    assert any(cell.value == -0.329 and cell.data_type == "n" for cell in cells)

    json_output = json.loads(export("json")["content"])
    json_strings = []
    def collect_strings(value):
        if isinstance(value, dict):
            for child in value.values():
                collect_strings(child)
        elif isinstance(value, list):
            for child in value:
                collect_strings(child)
        elif isinstance(value, str):
            json_strings.append(value)
    collect_strings(json_output)
    assert all(value in json_strings for value in dangerous.values())
    markdown = export("markdown")["content"]
    assert all(value in markdown for value in dangerous.values())
    assert displayed == original
    assert snapshot["calculation_fingerprint"] == source_fingerprint
    assert csv_output["calculation_fingerprint"] == source_fingerprint
    assert xlsx_output["calculation_fingerprint"] == source_fingerprint


def test_current_chapter18_xlsx_escapes_xml_illegal_provenance_without_rerunning(monkeypatch):
    workflow = workflow_for_method(METHOD_ID)
    displayed = workflow.starting_values(TEMPLATE_ID, "imperial")["displayed_inputs"]
    displayed["external_source_tool"] = "HCS\u000b7"
    snapshot = workflow.calculate(
        template_id=TEMPLATE_ID,
        unit_system="imperial",
        displayed_inputs=displayed,
    )
    original_displayed = deepcopy(displayed)
    original_normalized = deepcopy(snapshot["normalized_inputs"])
    original_result = deepcopy(snapshot["result"])
    original_calculation_fingerprint = snapshot["calculation_fingerprint"]
    original_snapshot_fingerprint = snapshot["input_snapshot_fingerprint"]
    monkeypatch.setattr(
        ch18.UrbanStreetSegmentMethod,
        "calculate",
        lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("export reran engine")),
    )

    def export(format_name):
        return export_current_workflow(
            METHOD_ID,
            template_id=TEMPLATE_ID,
            unit_system="imperial",
            displayed_inputs=displayed,
            calculation_fingerprint=original_calculation_fingerprint,
            input_snapshot_fingerprint=original_snapshot_fingerprint,
            result=snapshot["result"],
            export_format=format_name,
        )

    exported = export("xlsx")
    workbook = load_workbook(BytesIO(base64.b64decode(exported["content_base64"])))
    provenance = next(
        row[1]
        for row in workbook["Inputs"].iter_rows()
        if row[0].value == "external_source_tool"
    )

    assert provenance.value == r"HCS\u000B7"
    assert provenance.data_type == "s"
    assert "\u000b" not in provenance.value
    assert displayed == original_displayed
    assert snapshot["normalized_inputs"] == original_normalized
    assert snapshot["result"] == original_result
    assert exported["calculation_fingerprint"] == original_calculation_fingerprint
    assert snapshot["input_snapshot_fingerprint"] == original_snapshot_fingerprint
    assert exported["recalculated"] is False

    csv_export = export("csv")
    json_export = export("json")
    markdown_export = export("markdown")
    csv_cells = [cell for row in csv.reader(StringIO(csv_export["content"])) for cell in row]
    json_inputs = json.loads(json_export["content"])["inputs_summary"]

    raw_provenance = "HCS\u000b7"
    json_provenance = next(
        item["value"] for item in json_inputs if item["label"] == "external_source_tool"
    )
    assert raw_provenance in csv_cells
    assert json_provenance == raw_provenance
    assert raw_provenance in markdown_export["content"]
    assert r"HCS\u000B7" not in markdown_export["content"]
    assert displayed == original_displayed
    assert snapshot["normalized_inputs"] == original_normalized
    assert snapshot["result"] == original_result


def test_export_rejects_stale_inputs_and_tampered_chapter18_result_identity():
    snapshot = _chapter18_snapshot()
    changed = deepcopy(snapshot["displayed_inputs"])
    changed["curb_proportion"] = 0.71
    with pytest.raises(StaleResultError):
        export_current_workflow(METHOD_ID, template_id=TEMPLATE_ID, unit_system="imperial", displayed_inputs=changed,
            calculation_fingerprint=snapshot["calculation_fingerprint"], input_snapshot_fingerprint=snapshot["input_snapshot_fingerprint"], result=snapshot["result"], export_format="json")
    wrong = deepcopy(snapshot["result"])
    wrong["result_contract_version"] = "wrong"
    with pytest.raises(StaleResultError):
        export_current_workflow(METHOD_ID, template_id=TEMPLATE_ID, unit_system="imperial", displayed_inputs=snapshot["displayed_inputs"],
            calculation_fingerprint=snapshot["calculation_fingerprint"], input_snapshot_fingerprint=snapshot["input_snapshot_fingerprint"], result=wrong, export_format="json")
